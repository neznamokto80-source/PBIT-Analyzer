"""Интеграционные тесты на реальных .pbix.

Файлы не коммитятся: путь берётся из PBIT_TEST_PBIX_DIR (по умолчанию D:\\pbix).
Если папки или файлов нет, тесты помечаются skip, чтобы CI оставался зелёным.
"""
import os
from pathlib import Path

import openpyxl
import pytest
from pbi_modules.pbix_analyzer import PBIXAnalysisError, analyze_pbix_file

PBIX_DIR = Path(os.environ.get("PBIT_TEST_PBIX_DIR", r"D:\pbix"))


def _pbix_files():
    if not PBIX_DIR.is_dir():
        return []
    return sorted(PBIX_DIR.glob("*.pbix"))


pytestmark = pytest.mark.skipif(not _pbix_files(), reason=f"Нет .pbix в {PBIX_DIR}")


@pytest.fixture(scope="module")
def sample_pbix():
    return str(_pbix_files()[0])


def test_analyze_real_pbix_without_excel(tmp_path, sample_pbix):
    dfs = analyze_pbix_file(sample_pbix, str(tmp_path), create_excel=False)

    assert dfs["file_kind"] == "pbix"
    assert dfs["excel_path"] == ""
    assert not dfs["tables_ref"].empty
    assert not dfs["columns"].empty
    assert list(dfs["columns"].columns)[:4] == ["№ столбца", "№ таблицы", "Таблица", "Столбец"]
    assert not dfs["lineage"].empty
    assert not dfs["stats"].empty
    assert "Тип файла" in set(dfs["stats"]["Параметр"])
    assert dfs["warnings"] == []


def test_analyze_real_pbix_builds_excel(tmp_path, sample_pbix):
    dfs = analyze_pbix_file(sample_pbix, str(tmp_path), create_excel=True)

    excel_path = Path(dfs["excel_path"])
    assert excel_path.is_file()
    workbook = openpyxl.load_workbook(excel_path)
    assert workbook.sheetnames[0] == "Справочник таблиц"
    assert "Безопасность (RLS)" in workbook.sheetnames
    assert "Статистика хранения" in workbook.sheetnames
    assert "DAX" not in workbook.sheetnames
    assert "Power Query" not in workbook.sheetnames
    assert "Перспективы" not in workbook.sheetnames
    assert "Визуальные элементы" not in workbook.sheetnames
    assert "Листы" not in workbook.sheetnames


def test_lineage_focus_on_real_pbix(tmp_path, sample_pbix):
    dfs = analyze_pbix_file(sample_pbix, str(tmp_path), create_excel=False, lineage_focus="s1")

    assert not dfs["lineage_upstream"].empty
    assert set(dfs["lineage_upstream"]["ID фокуса"].unique()) == {"s1"}


def test_broken_pbix_raises_clear_error(tmp_path):
    broken = tmp_path / "broken.pbix"
    broken.write_bytes(b"not a zip at all")

    with pytest.raises(PBIXAnalysisError, match="Не удалось прочитать PBIX"):
        analyze_pbix_file(str(broken), str(tmp_path), create_excel=False)


def test_on_disk_mode_cleans_temp_dir(tmp_path, sample_pbix):
    dfs = analyze_pbix_file(sample_pbix, str(tmp_path), create_excel=False, on_disk=True)

    assert dfs["model_size"] >= 0
    assert not list(tmp_path.glob("_temp_pbix_*"))
